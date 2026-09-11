"""Parse SealDice import worksheets without modifying the source workbook."""

from __future__ import annotations

import io
import os
import re
import zipfile
from typing import Any


COC_ALIASES = {
    "str": "力量", "力量": "力量",
    "dex": "敏捷", "敏捷": "敏捷",
    "con": "体质", "体质": "体质", "體質": "体质",
    "siz": "体型", "体型": "体型", "體型": "体型",
    "app": "外貌", "外貌": "外貌", "外表": "外貌",
    "int": "智力", "智力": "智力", "灵感": "智力", "靈感": "智力",
    "pow": "意志", "意志": "意志",
    "edu": "教育", "教育": "教育",
    "luck": "幸运", "幸运": "幸运", "运气": "幸运",
    "san": "理智", "san值": "理智", "理智": "理智", "理智值": "理智",
    "hp": "生命值", "生命值": "生命值", "生命": "生命值", "体力": "生命值",
    "mp": "魔法值", "魔法值": "魔法值", "魔法": "魔法值",
    "mov": "移动力", "移动力": "移动力",
    "armor": "护甲", "护甲": "护甲", "装甲": "护甲",
    "build": "体格", "体格": "体格",
    "信用": "信用评级", "信誉": "信用评级", "信用评级": "信用评级",
    "图书馆": "图书馆使用", "图书馆使用": "图书馆使用",
    "计算机": "计算机使用", "电脑": "计算机使用", "计算机使用": "计算机使用",
    "开锁": "锁匠", "撬锁": "锁匠", "锁匠": "锁匠",
    "汽车": "汽车驾驶", "驾驶": "汽车驾驶", "汽车驾驶": "汽车驾驶",
    "领航": "导航", "导航": "导航",
    "重型": "操作重型机械", "重型机械": "操作重型机械", "重型操作": "操作重型机械",
    "步枪/霰弹枪": "射击:步霰", "步枪": "射击:步霰", "霰弹枪": "射击:步霰",
    "射击:步枪/霰弹枪": "射击:步霰", "射击:步枪": "射击:步霰", "射击:霰弹枪": "射击:步霰",
    "步槍/霰彈槍": "射击:步霰", "射擊:步槍/霰彈槍": "射击:步霰",
    "手枪": "射击:手枪", "机枪": "射击:机枪", "冲锋枪": "射击:冲锋枪",
    "克苏鲁": "克苏鲁神话", "cm": "克苏鲁神话", "克苏鲁神话": "克苏鲁神话",
    "斗殴": "斗殴", "格斗": "格斗",
}

DND_ALIASES = {
    "str": "力量", "strength": "力量", "力量": "力量",
    "dex": "敏捷", "dexterity": "敏捷", "敏捷": "敏捷",
    "con": "体质", "constitution": "体质", "体质": "体质", "體質": "体质",
    "int": "智力", "intelligence": "智力", "智力": "智力",
    "wis": "感知", "wisdom": "感知", "感知": "感知",
    "cha": "魅力", "charisma": "魅力", "魅力": "魅力",
    "ac": "ac", "护甲": "ac", "护甲等级": "ac",
    "hp": "hp", "生命值": "hp", "生命": "hp",
    "hpmax": "hpmax", "生命值上限": "hpmax", "生命上限": "hpmax",
    "dc": "dc", "法术豁免": "dc", "pp": "pp", "被动察觉": "pp",
    "熟练": "熟练", "熟练加值": "熟练",
}

DND_SKILL_PARENT = {
    "运动": "力量",
    "体操": "敏捷", "巧手": "敏捷", "隐匿": "敏捷",
    "调查": "智力", "奥秘": "智力", "历史": "智力", "自然": "智力", "宗教": "智力",
    "察觉": "感知", "洞悉": "感知", "驯兽": "感知", "医药": "感知", "求生": "感知",
    "游说": "魅力", "欺瞒": "魅力", "威吓": "魅力", "表演": "魅力",
}

DND_KNOWN_NAMES = set(DND_ALIASES.values()) | set(DND_SKILL_PARENT) | {
    "力量豁免", "敏捷豁免", "体质豁免", "智力豁免", "感知豁免", "魅力豁免",
}

PLACEHOLDER_WORDS = (
    "请复制", "快捷输入", "不知道用什么导入", "骰娘指令", "电脑文本输入框",
    "请在这里", "请选择类型", "原技能名称", "当前时间", "累计消耗",
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    value = str(value).replace("\r", " ").replace("\n", " ")
    value = re.sub(r"[\x00-\x1f\x7f]", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _number(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if value != value:
            return None
        return int(value) if int(value) == value else float(value)
    text = _text(value).replace("％", "%")
    match = re.fullmatch(r"[-+]?\d+(?:\.\d+)?%?", text)
    if not match:
        return None
    number = float(text.rstrip("%"))
    return int(number) if number.is_integer() else number


def _clean_label(value: Any) -> str:
    text = _text(value)
    text = re.sub(r"^[.。·•\s]+|[：:]$", "", text)
    text = text.replace("Ω", "").strip()
    return text[:40]


def _canonical(label: str, system: str) -> str:
    clean = _clean_label(label)
    key = clean.replace("：", ":").lower()
    aliases = DND_ALIASES if system == "dnd" else COC_ALIASES
    canonical = aliases.get(key, aliases.get(clean, clean))
    # DiceScript treats ASCII slash in an st attribute name as division.
    # Known COC names are mapped above; preserve unknown names with a fullwidth slash.
    return canonical.replace("/", "／")


def _valid_entry(text: str, max_len: int = 80) -> bool:
    if not text or len(text) > max_len or text in {"0", "无", "暂无", "——", "--", "×", "√", "☐", "☑"}:
        return False
    return not any(word in text for word in PLACEHOLDER_WORDS)


def _strip_bad_validations(raw_bytes: bytes) -> io.BytesIO:
    src = io.BytesIO(raw_bytes)
    dst = io.BytesIO()
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item.startswith("xl/worksheets/") and item.endswith(".xml"):
                xml = data.decode("utf-8", errors="ignore")
                xml = re.sub(r"<(?:\w+:)?dataValidations\b[\s\S]*?</(?:\w+:)?dataValidations>", "", xml)
                xml = re.sub(r"<extLst>[\s\S]*?</extLst>", "", xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    dst.seek(0)
    return dst


def _system_score(names: list[str], sheets: dict[str, list[list[Any]]] | None = None) -> tuple[int, int]:
    joined_names = " ".join(_text(name).lower() for name in names)
    dnd_score = sum(3 for marker in ("dnd", "5e", "龙与地下城") if marker in joined_names)
    coc_score = sum(3 for marker in ("coc", "克苏鲁", "调查员", "简化卡") if marker in joined_names)
    dnd_score += sum(2 for marker in ("起源", "主要", "背包", "施法", "法术书") if marker in joined_names)

    if sheets:
        texts = []
        for name, matrix in sheets.items():
            if "法术大全" in name:
                continue
            texts.extend(_text(value).lower() for row in matrix[:100] for value in row[:80] if value is not None)
        content = " ".join(texts)
        if re.search(r"\.(?:dst|cst)\b", content, re.IGNORECASE):
            dnd_score += 20
        dnd_score += sum(4 for marker in ("hpmax", "法术位", "法术豁免", "法术攻击", "熟练") if marker in content)
        dnd_abilities = sum(marker in content for marker in ("力量", "敏捷", "体质", "智力", "感知", "魅力"))
        if dnd_abilities >= 5:
            dnd_score += 12

        coc_score += sum(5 for marker in ("体型", "意志", "教育", "幸运", "理智", "克苏鲁神话") if marker in content)
        coc_score += sum(3 for marker in ("siz", "pow", "edu", "luck", "san") if re.search(rf"\b{marker}\b", content, re.IGNORECASE))
    return dnd_score, coc_score


def _choose_import_sheet(
    names: list[str],
    card_system: str,
    sheets: dict[str, list[list[Any]]] | None = None,
) -> tuple[str, str]:
    candidates = [name for name in names if "导入" in _text(name)]

    dnd_markers = ("dnd", "5e", "龙与地下城")
    coc_markers = ("coc", "克苏鲁", "调查员", "简化卡")
    detected = card_system
    if detected == "auto":
        dnd_score, coc_score = _system_score(names, sheets)
        detected = "dnd" if dnd_score > coc_score else "coc"

    if not candidates:
        if detected == "dnd":
            for fallback in ("主要情况", "主要", "施法", "角色", "起源"):
                match = next((name for name in names if _text(name) == fallback), None)
                if match:
                    return match, detected
        raise ValueError("工作簿中没有找到名称含“导入”的工作表")

    def score(name: str) -> int:
        low = _text(name).lower()
        wanted = dnd_markers if detected == "dnd" else coc_markers
        unwanted = coc_markers if detected == "dnd" else dnd_markers
        value = 20
        value += sum(8 for marker in wanted if marker in low)
        value -= sum(12 for marker in unwanted if marker in low)
        if "骰娘" in low:
            value += 4
        return value

    return max(candidates, key=score), detected


def _load_import_matrix(
    file_content: bytes,
    filename: str,
    card_system: str,
) -> tuple[str, str, list[list[Any]], dict[str, list[list[Any]]]]:
    extension = os.path.splitext(filename)[1].lower()
    if file_content.startswith(b"PK\x03\x04"):
        extension = ".xlsx"
    elif file_content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        extension = ".xls"

    if extension in (".xlsx", ".xlsm"):
        import openpyxl

        def read_relevant(source: io.BytesIO) -> tuple[list[str], dict[str, list[list[Any]]]]:
            workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
            try:
                names = list(workbook.sheetnames)
                relevant = {
                    name for name in names
                    if "导入" in _text(name) or "人物" in _text(name) or "角色" in _text(name) or _text(name) in {
                        "起源", "角色", "主要", "主要情况", "背包", "装备", "施法", "法术书", "法术大全"
                    }
                }
                matrices = {
                    name: [list(row) for row in workbook[name].iter_rows(values_only=True)]
                    for name in relevant
                }
                return names, matrices
            finally:
                workbook.close()

        try:
            names, support_sheets = read_relevant(io.BytesIO(file_content))
        except Exception:
            names, support_sheets = read_relevant(_strip_bad_validations(file_content))
        sheet_name, detected = _choose_import_sheet(names, card_system, support_sheets)
        matrix = support_sheets[sheet_name]
        return sheet_name, detected, matrix, support_sheets

    if extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=file_content)
        names = [sheet.name for sheet in workbook.sheets()]
        relevant = {
            name for name in names
            if "导入" in _text(name) or "人物" in _text(name) or "角色" in _text(name) or _text(name) in {
                "起源", "角色", "主要", "主要情况", "背包", "装备", "施法", "法术书", "法术大全"
            }
        }
        support_sheets = {}
        for name in relevant:
            sheet = workbook.sheet_by_name(name)
            support_sheets[name] = [sheet.row_values(row) for row in range(sheet.nrows)]
        sheet_name, detected = _choose_import_sheet(names, card_system, support_sheets)
        return sheet_name, detected, support_sheets[sheet_name], support_sheets

    raise ValueError("一键导入仅支持 .xlsx、.xlsm 和 .xls 文件")


def _cell(matrix: list[list[Any]], row: int, column: int) -> Any:
    if row < 0 or row >= len(matrix) or column < 0 or column >= len(matrix[row]):
        return None
    return matrix[row][column]


def _command_cells(matrix: list[list[Any]]) -> list[str]:
    commands = []
    for row in matrix:
        for value in row:
            text = _text(value)
            if re.match(r"^\.(?:st|dst|cst)\b", text, re.IGNORECASE):
                commands.append(text)
    return commands


def _extract_name(matrix: list[list[Any]], commands: list[str]) -> str:
    for row in matrix:
        for value in row:
            text = _text(value)
            match = re.match(r"^\.nn\s+(.{1,40})$", text, re.IGNORECASE)
            if match:
                return _clean_label(match.group(1))

    for command in sorted(commands, key=len, reverse=True):
        body = re.sub(r"^\.(?:st|dst|cst)\s*", "", command, flags=re.IGNORECASE)
        match = re.match(r"([^\d:：*?？-]{1,40})[-?？](?=[^\s]*\d)", body)
        if match:
            return _clean_label(match.group(1))

    for row in matrix[:8]:
        for value in row[:5]:
            text = _text(value)
            if any(mark in text for mark in ("，", ",")) and not any(word in text for word in PLACEHOLDER_WORDS):
                candidate = re.split(r"[，,]", text, maxsplit=1)[0].strip()
                if 1 <= len(candidate) <= 40 and not re.search(r"\d", candidate):
                    return candidate

    rejected_labels = {"注释", "备注", "说明", "年龄", "性别", "职业", "玩家", "姓名"}
    for row in matrix:
        for index, value in enumerate(row):
            if _clean_label(value).lower() in {"姓名", "name", "角色名", "人物名"}:
                for candidate in row[index + 1:index + 6]:
                    text = _text(candidate)
                    if _valid_entry(text, 40) and _clean_label(text) not in rejected_labels:
                        return text
    return ""


def _extract_labeled_text(matrix: list[list[Any]], labels: tuple[str, ...], limit: int = 600) -> str:
    """Read a nearby free-text value such as appearance from varied card layouts."""
    wanted = {str(label).strip().lower() for label in labels}
    reject = wanted | {"力量", "敏捷", "体质", "体型", "智力", "意志", "教育", "幸运", "理智", "生命值", "魔法值", "姓名", "年龄", "性别", "职业", "玩家", "str", "dex", "con", "siz", "app", "int", "pow", "edu", "luck", "san", "hp", "mp"}
    for row_index, row in enumerate(matrix):
        for column, value in enumerate(row):
            if _clean_label(value).lower() not in wanted:
                continue
            candidates = list(row[column + 1:column + 5])
            if row_index + 1 < len(matrix):
                candidates.extend(matrix[row_index + 1][max(0, column - 1):column + 4])
            for candidate in candidates:
                text = _text(candidate)
                if text and text.lower() not in reject and not re.fullmatch(r'-?\d+(?:\.\d+)?', text) and not any(word in text for word in PLACEHOLDER_WORDS):
                    return text[:limit]
    return ""


def _parse_spaced_command(command: str, system: str) -> tuple[dict[str, int | float], dict[str, float]]:
    body = re.sub(r"^\.(?:st|dst|cst)\s*", "", command, flags=re.IGNORECASE)
    result: dict[str, int | float] = {}
    proficiency: dict[str, float] = {}
    pattern = re.compile(r"([^\s:：*?？-]{1,30})(\*(?:\d+(?:\.\d+)?)?)?\s*[:：]\s*(-?\d+(?:\.\d+)?)")
    for match in pattern.finditer(body):
        label = _canonical(match.group(1), system)
        if not label or any(word in label for word in PLACEHOLDER_WORDS):
            continue
        number = _number(match.group(3))
        if number is None:
            continue
        result[label] = number
        if match.group(2):
            factor_text = match.group(2)[1:]
            proficiency[label] = float(factor_text) if factor_text else 1.0
    return result, proficiency


def _extract_coc_grid(matrix: list[list[Any]]) -> tuple[dict[str, int | float], dict[str, int | float]]:
    attributes: dict[str, int | float] = {}
    skills: dict[str, int | float] = {}
    instruction_row = next((i for i, row in enumerate(matrix) if any("快捷输入" in _text(v) for v in row)), min(60, len(matrix)))

    for row_index in range(min(instruction_row, len(matrix))):
        for label_column in (1, 4, 7, 10):
            label = _clean_label(_cell(matrix, row_index, label_column))
            value = _number(_cell(matrix, row_index, label_column + 1))
            canonical = _canonical(label, "coc")
            if canonical in {"力量", "敏捷", "体质", "体型", "外貌", "智力", "意志", "教育", "幸运", "理智", "生命值", "魔法值", "移动力", "体格"} and value is not None:
                attributes[canonical] = value

        for label_column, value_column in ((13, 16), (19, 22)):
            label = _clean_label(_cell(matrix, row_index, label_column))
            extra = _clean_label(_cell(matrix, row_index, label_column + 1))
            value = _number(_cell(matrix, row_index, value_column))
            if not label or value is None or value <= 0:
                continue
            base = label.rstrip("：:")
            if extra and base in {"技艺", "艺术与手艺", "外语", "科学", "生存", "格斗", "射击", "驾驶"}:
                if base == "射击":
                    label = f"射击:{extra}"
                elif base == "驾驶":
                    label = f"驾驶:{extra}"
                else:
                    label = extra
            canonical = _canonical(label, "coc")
            if _valid_entry(canonical, 30):
                skills[canonical] = value
    return attributes, skills


def _finalize_dnd_values(raw: dict[str, int | float], proficiency: dict[str, float]) -> tuple[dict[str, int | float], dict[str, int | float]]:
    attributes: dict[str, int | float] = {}
    skills: dict[str, int | float] = {}
    ability_names = {"力量", "敏捷", "体质", "智力", "感知", "魅力"}
    basic_names = ability_names | {"ac", "hp", "hpmax", "dc", "hd", "pp", "熟练", "体型"}
    for name, value in raw.items():
        if name in basic_names:
            attributes[name] = value

    proficiency_bonus = float(raw.get("熟练", 0) or 0)
    for name, parent in DND_SKILL_PARENT.items():
        if name not in raw:
            continue
        ability = float(raw.get(parent, 10) or 10)
        base = float(raw[name])
        factor = proficiency.get(name, 0)
        total = base + int(ability // 2) - 5 + proficiency_bonus * factor
        skills[name] = int(total) if float(total).is_integer() else total
    for name, value in raw.items():
        if name not in basic_names and name not in skills and _valid_entry(name, 30):
            skills[name] = value
    return attributes, skills


def _extract_dnd_grid(matrix: list[list[Any]]) -> dict[str, int | float]:
    raw: dict[str, int | float] = {}
    for row_index, row in enumerate(matrix):
        for column, value in enumerate(row):
            label_text = _text(value)
            canonical = _canonical(label_text, "dnd")
            if canonical not in DND_KNOWN_NAMES and " " in label_text:
                canonical = _canonical(label_text.split(" ", 1)[0], "dnd")
            if canonical not in DND_KNOWN_NAMES:
                continue
            candidates = [
                _cell(matrix, row_index, column + 1),
                _cell(matrix, row_index, column + 2),
                _cell(matrix, row_index, column + 3),
                _cell(matrix, row_index, column + 4),
                _cell(matrix, row_index + 1, column),
            ]
            number = next((_number(candidate) for candidate in candidates if _number(candidate) is not None), None)
            if number is not None:
                raw[canonical] = number
    return raw


def _extract_dnd_weapons(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    weapons = []
    placeholder_names = {
        "无", "名称", "武器名称", "武器", "左手", "右手", "双手", "其他", "头部", "颈部",
        "肩部", "背部", "胸部", "腕部", "手部", "腰部", "腿部", "脚部", "装备", "徒手打击",
    }
    for header_row, row in enumerate(matrix):
        labels = [_clean_label(value) for value in row]
        section_columns = [column for column, label in enumerate(labels) if label == "武器"]
        damage_columns = [column for column, label in enumerate(labels) if label in {"伤害骰", "武器伤害", "伤害"}]
        if not section_columns or not damage_columns:
            continue
        section_column = section_columns[0]
        damage_column = next((column for column in damage_columns if column > section_column), None)
        if damage_column is None:
            continue
        name_column = next((
            column for column, label in enumerate(labels)
            if column > section_column and label in {"名称", "武器名称"}
        ), None)
        property_column = next((column for column, label in enumerate(labels) if column > section_column and label == "特性"), damage_column)
        type_column = next((column for column, label in enumerate(labels) if column > damage_column and label in {"类型", "伤害类型"}), None)
        ammo_column = next((column for column, label in enumerate(labels) if column > damage_column and label in {"弹药", "弹药数"}), None)

        for data_row in range(header_row + 1, min(len(matrix), header_row + 18)):
            dice_column = None
            damage = ""
            for column in range(max(section_column, damage_column - 2), min(len(matrix[data_row]), damage_column + 5)):
                text = _text(_cell(matrix, data_row, column))
                if re.search(r"\d+d\d+", text, re.IGNORECASE):
                    dice_column = column
                    damage = text
                    next_text = _text(_cell(matrix, data_row, column + 1))
                    if re.fullmatch(r"[+-]\s*\d+", next_text):
                        damage += next_text.replace(" ", "")
                    break
            if dice_column is None:
                continue

            name = _clean_label(_cell(matrix, data_row, name_column)) if name_column is not None else ""
            section_name = _clean_label(_cell(matrix, data_row, section_column))
            if (not _valid_entry(name, 60) or name in placeholder_names) and _valid_entry(section_name, 60) and section_name not in placeholder_names:
                name = section_name
            if not _valid_entry(name, 60) or name in placeholder_names:
                continue
            properties = _text(_cell(matrix, data_row, property_column))
            attack_modifier = None
            for column in range(property_column + 1, dice_column):
                text = _text(_cell(matrix, data_row, column))
                match = re.fullmatch(r"\+\s*(\d+)", text)
                number = _number(text)
                if match:
                    attack_modifier = int(match.group(1))
                elif number is not None and -20 <= number <= 30:
                    attack_modifier = int(number)
            damage_type = _text(_cell(matrix, data_row, type_column)) if type_column is not None else ""
            ammo = 0
            ammo_type = ""
            if ammo_column is not None:
                for column in range(ammo_column, min(len(matrix[data_row]), ammo_column + 4)):
                    value = _cell(matrix, data_row, column)
                    number = _number(value)
                    if number is not None and 0 <= number <= 9999:
                        ammo = int(number)
                    elif _valid_entry(_text(value), 30):
                        ammo_type = _text(value)
            weapon = {
                "name": name,
                "type": damage_type,
                "skill": "",
                "skill_value": 0,
                "attack_modifier": attack_modifier,
                "damage": damage,
                "range": properties,
                "attacks": "",
                "capacity": ammo,
                "ammo": ammo,
                "ammo_type": ammo_type,
                "malfunction": None,
            }
            if name not in {item["name"] for item in weapons}:
                weapons.append(weapon)
        if weapons:
            break
    return weapons[:20]


def _extract_weapons(matrix: list[list[Any]], system: str, skills: dict[str, Any]) -> list[dict[str, Any]]:
    if system == "dnd":
        return _extract_dnd_weapons(matrix)
    header = None
    for row_index, row in enumerate(matrix):
        for column, value in enumerate(row):
            if _clean_label(value) in {"武器", "武器表", "武器列表"}:
                header = (row_index, column)
                break
        if header:
            break
    if not header:
        return []

    weapons = []
    start_row, start_column = header
    stop_words = {"资产", "随身物品", "携带物品", "道具", "背景故事", "背景"}
    for row_index in range(start_row + 1, min(len(matrix), start_row + 24)):
        name = _text(_cell(matrix, row_index, start_column))
        if _clean_label(name) in stop_words:
            break
        if not _valid_entry(name, 60) or name in {"武器表", "请选择类型"}:
            continue
        row_values = [_cell(matrix, row_index, start_column + offset) for offset in range(12)]
        if system == "coc" and len(row_values) >= 12:
            skill = _canonical(_text(row_values[2]), "coc")
            chance = _number(row_values[3])
            capacity = _number(row_values[10])
            capacity = int(capacity) if capacity is not None and capacity > 0 else 0
            weapon = {
                "name": name,
                "type": _text(row_values[1]),
                "skill": skill,
                "skill_value": skills.get(skill, chance or 0),
                "damage": _text(row_values[6]),
                "range": _text(row_values[7]),
                "attacks": _text(row_values[9]),
                "capacity": capacity,
                "ammo": capacity,
                "malfunction": _number(row_values[11]),
            }
        else:
            texts = [_text(value) for value in row_values if _valid_entry(_text(value), 60)]
            damage = next((value for value in texts[1:] if re.search(r"\d+d\d+", value, re.IGNORECASE)), "")
            weapon = {
                "name": name,
                "type": texts[1] if len(texts) > 1 else "",
                "skill": texts[2] if len(texts) > 2 else "",
                "skill_value": 0,
                "damage": damage,
                "range": "",
                "attacks": "",
                "capacity": 0,
                "ammo": 0,
                "malfunction": None,
            }
        if weapon["name"] not in {item["name"] for item in weapons}:
            weapons.append(weapon)
    return weapons[:20]


def _extract_dnd_items(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    headings = {"名称", "冒险装备", "工具", "道具", "物品", "其他物品", "背包"}
    rejected = headings | {
        "描述", "价格", "重量", "磅", "数量", "合计结余", "收入项目", "支出项目",
        "背包2", "次元袋", "项目", "明细",
    }
    items = []
    tables = []
    for row_index, row in enumerate(matrix):
        header_columns = [column for column, value in enumerate(row) if _clean_label(value) in headings]
        for position, name_column in enumerate(header_columns):
            next_column = header_columns[position + 1] if position + 1 < len(header_columns) else len(row)
            quantity_column = next((
                column for column in range(name_column + 1, next_column)
                if _clean_label(_cell(matrix, row_index, column)) == "数量"
            ), None)
            tables.append((row_index, name_column, quantity_column))

    for header_row, name_column, quantity_column in tables:
        for data_row in range(header_row + 1, min(len(matrix), header_row + 14)):
            name = _clean_label(_cell(matrix, data_row, name_column))
            if name in headings:
                break
            if not _valid_entry(name, 80) or name in rejected or re.match(r"^[x×><+-]$", name, re.IGNORECASE):
                continue
            if _number(name) is not None or re.match(r"^(?:GP|SP|CP|EP|PP)?\s*\d", name, re.IGNORECASE):
                continue
            quantity_value = _number(_cell(matrix, data_row, quantity_column)) if quantity_column is not None else None
            quantity = int(quantity_value) if quantity_value is not None and 0 < quantity_value <= 999 else 1
            if name not in {item["name"] for item in items}:
                items.append({"name": name, "quantity": quantity})
    return items[:100]


def _extract_items(matrix: list[list[Any]], system: str = "coc") -> list[dict[str, Any]]:
    if system == "dnd":
        return _extract_dnd_items(matrix)
    header = None
    headings = {"随身物品", "携带物品", "道具", "装备", "背包"}
    for row_index, row in enumerate(matrix):
        for column, value in enumerate(row):
            if _clean_label(value) in headings:
                header = (row_index, column)
                break
        if header:
            break
    if not header:
        return []

    items = []
    start_row, start_column = header
    stop_words = {"背景", "背景故事", "描述", "信仰", "重要人", "重要地", "宝物", "特质", "伤疤", "恐惧"}
    for row_index in range(start_row + 1, min(len(matrix), start_row + 24)):
        row = matrix[row_index]
        if any(_clean_label(value) in stop_words for value in row[:start_column + 3]):
            break
        for column in range(start_column, min(len(row), start_column + 12)):
            name = _text(row[column])
            if not _valid_entry(name, 80) or _number(row[column]) is not None:
                continue
            if re.match(r"^\.(?:st|dst|cst|nn)\b", name, re.IGNORECASE):
                continue
            if _clean_label(name) in headings or any(word in name for word in PLACEHOLDER_WORDS):
                continue
            quantity_value = _number(row[column + 1]) if column + 1 < len(row) else None
            quantity = int(quantity_value) if quantity_value is not None and 0 < quantity_value <= 999 else 1
            if name not in {item["name"] for item in items}:
                items.append({"name": name, "quantity": quantity})
    return items[:100]


def _find_nearby_number(
    matrix: list[list[Any]],
    row_index: int,
    column: int,
    max_distance: int = 6,
) -> int | float | None:
    candidates = []
    for distance in range(1, max_distance + 1):
        candidates.extend((
            _cell(matrix, row_index, column + distance),
            _cell(matrix, row_index + 1, column + distance - 1),
            _cell(matrix, row_index + 1, column),
        ))
    return next((_number(value) for value in candidates if _number(value) is not None), None)


def _extract_spellcasting_numbers(
    sheets: dict[str, list[list[Any]]],
    attributes: dict[str, Any],
) -> tuple[int | None, int | None]:
    save_dcs = []
    attack_bonuses = []
    for sheet_name in ("施法", "主要", "主要情况"):
        matrix = sheets.get(sheet_name)
        if not matrix:
            continue
        for row_index, row in enumerate(matrix):
            for column, value in enumerate(row):
                label = _clean_label(value).replace(" ", "")
                if label in {"法术豁免DC", "法术DC"}:
                    number = _find_nearby_number(matrix, row_index, column)
                    if number is not None and 5 <= number <= 40:
                        save_dcs.append(int(number))
                if label in {"法术命中", "法术攻击", "法术攻击骰"}:
                    candidates = [_cell(matrix, row_index, column + offset) for offset in range(1, 8)]
                    candidates.extend(_cell(matrix, row_index + offset, column) for offset in range(1, 6))
                    for candidate in candidates:
                        text = _text(candidate)
                        match = re.search(r"(?:d20\s*)?([+-]\s*\d+)", text, re.IGNORECASE)
                        number = _number(candidate)
                        if match:
                            attack_bonuses.append(int(match.group(1).replace(" ", "")))
                            break
                        if number is not None and -20 <= number <= 30:
                            attack_bonuses.append(int(number))
                            break

    attr_dc = _number(attributes.get("dc"))
    save_dc = int(attr_dc) if attr_dc is not None else (max(save_dcs) if save_dcs else None)
    attack_bonus = max(attack_bonuses) if attack_bonuses else (save_dc - 8 if save_dc is not None else None)
    return save_dc, attack_bonus


def _extract_spell_slots(sheets: dict[str, list[list[Any]]]) -> dict[str, dict[str, int]]:
    slots: dict[str, dict[str, int]] = {}
    for sheet_name in ("施法", "主要", "主要情况"):
        matrix = sheets.get(sheet_name)
        if not matrix:
            continue
        for row_index, row in enumerate(matrix):
            level_columns = [column for column, value in enumerate(row) if _clean_label(value) == "环阶"]
            slot_columns = [column for column, value in enumerate(row) if _clean_label(value) in {"法术位", "环位"}]
            for level_column in level_columns:
                slot_column = next((column for column in slot_columns if column > level_column), None)
                if slot_column is None:
                    continue
                for data_row in range(row_index + 1, min(len(matrix), row_index + 12)):
                    level = _number(_cell(matrix, data_row, level_column))
                    if level is None or not 0 <= level <= 9:
                        break
                    level = int(level)
                    if level == 0:
                        continue
                    numeric = []
                    for column in range(slot_column, min(slot_column + 5, len(matrix[data_row]))):
                        number = _number(_cell(matrix, data_row, column))
                        if number is not None and 0 <= number <= 99:
                            numeric.append(int(number))
                    if not numeric:
                        continue
                    current = numeric[0]
                    maximum = numeric[-1]
                    if maximum > 0:
                        slots[str(level)] = {"current": min(current, maximum), "max": maximum}
                if slots:
                    return slots
    return slots


def _spell_detail_index(matrix: list[list[Any]]) -> dict[str, dict[str, Any]]:
    for row_index, row in enumerate(matrix[:10]):
        headers = {_clean_label(value): column for column, value in enumerate(row) if _clean_label(value)}
        name_column = headers.get("法术名", headers.get("法术名称"))
        level_column = headers.get("环阶", headers.get("LV"))
        description_column = headers.get("法术详述", headers.get("法术效果"))
        if name_column is None or level_column is None or description_column is None:
            continue
        details = {}
        for data_row in range(row_index + 1, len(matrix)):
            name = _clean_label(_cell(matrix, data_row, name_column))
            level = _number(_cell(matrix, data_row, level_column))
            description = _text(_cell(matrix, data_row, description_column))
            if _valid_entry(name, 60) and level is not None and 0 <= level <= 9:
                details[name] = {"level": int(level), "description": description}
        return details
    return {}


def _extract_selected_spells(sheets: dict[str, list[list[Any]]]) -> list[dict[str, Any]]:
    selected: dict[str, dict[str, Any]] = {}
    spellbook = sheets.get("法术书")
    if spellbook:
        for row_index, row in enumerate(spellbook):
            for column, value in enumerate(row):
                if "法术列表" not in _text(value):
                    continue
                for data_row in range(row_index + 1, len(spellbook)):
                    level = _number(_cell(spellbook, data_row, column + 1))
                    name = _clean_label(_cell(spellbook, data_row, column + 3))
                    if level is None and not name:
                        if data_row > row_index + 3:
                            break
                        continue
                    if level is not None and 0 <= level <= 9 and _valid_entry(name, 60):
                        selected[name] = {"name": name, "level": int(level), "description": ""}
                break

    for sheet_name in ("施法", "主要", "主要情况"):
        matrix = sheets.get(sheet_name)
        if not matrix:
            continue
        for row_index, row in enumerate(matrix):
            for name_column, value in enumerate(row):
                if _clean_label(value) not in {"法术名", "法术名称"}:
                    continue
                level_column = next((
                    column for column in range(name_column - 1, max(-1, name_column - 10), -1)
                    if _clean_label(_cell(matrix, row_index, column)) in {"LV", "Lv", "环阶"}
                ), None)
                if level_column is None:
                    continue
                description_column = next((
                    column for column in range(name_column + 1, min(len(row), name_column + 35))
                    if _clean_label(_cell(matrix, row_index, column)) in {"法术效果", "法术详述"}
                ), None)
                for data_row in range(row_index + 1, min(len(matrix), row_index + 16)):
                    name = _clean_label(_cell(matrix, data_row, name_column))
                    level = _number(_cell(matrix, data_row, level_column))
                    if _clean_label(name) in {"法术名", "法术名称"}:
                        break
                    if level is not None and 0 <= level <= 9 and _valid_entry(name, 60):
                        description = _text(_cell(matrix, data_row, description_column)) if description_column is not None else ""
                        selected[name] = {"name": name, "level": int(level), "description": description}

    details = _spell_detail_index(sheets.get("法术大全", []))
    for name, spell in selected.items():
        detail = details.get(name)
        if detail:
            spell["level"] = detail["level"]
            if detail["description"]:
                spell["description"] = detail["description"]
    return list(selected.values())[:200]


def _classify_spell(spell: dict[str, Any]) -> dict[str, Any]:
    description = _text(spell.get("description"))
    save_match = re.search(r"(力量|敏捷|体质|智力|感知|魅力)(?:豁免|豁免检定)", description)
    attack = bool(re.search(r"(?:近战|远程)?法术攻击(?:检定)?|以法术发动.{0,8}攻击", description))
    healing = bool(re.search(r"(?:恢复|回复|治疗).{0,20}(?:生命值|HP)|生命值.{0,12}(?:恢复|回复)", description, re.IGNORECASE))
    dice_match = re.search(r"(\d+d\d+(?:\s*[+-]\s*\d+)?)", description, re.IGNORECASE)
    upcast_match = re.search(r"(?:升环施法|使用更高环阶).{0,160}?(\d+d\d+)", description, re.IGNORECASE)
    if save_match:
        roll_type = "save"
    elif attack:
        roll_type = "attack"
    elif healing:
        roll_type = "healing"
    elif dice_match:
        roll_type = "damage"
    else:
        roll_type = "none"
    spell.update({
        "roll_type": roll_type,
        "save_ability": save_match.group(1) if save_match else "",
        "damage_dice": dice_match.group(1).replace(" ", "") if dice_match else "",
        "upcast_dice": upcast_match.group(1) if upcast_match else "",
    })
    return spell


def parse_import_matrix(
    matrix: list[list[Any]],
    system: str,
    sheet_name: str = "导入",
    support_sheets: dict[str, list[list[Any]]] | None = None,
) -> dict[str, Any]:
    commands = _command_cells(matrix)
    name = _extract_name(matrix, commands)
    proficiencies: dict[str, float] = {}

    if system == "coc":
        attributes, skills = _extract_coc_grid(matrix)
        for command in commands:
            raw, _ = _parse_spaced_command(command, system)
            for key, value in raw.items():
                if key in attributes:
                    attributes[key] = value
                elif key in skills:
                    skills[key] = value
        st_entries = [
            {"name": key, "value": value, "proficiency": 0}
            for key, value in {**attributes, **skills}.items()
        ]
    else:
        raw: dict[str, int | float] = {}
        for command in commands:
            parsed, prof = _parse_spaced_command(command, system)
            if len(parsed) > len(raw):
                raw, proficiencies = parsed, prof
        if not raw:
            raw = _extract_dnd_grid(matrix)
        attributes, skills = _finalize_dnd_values(raw, proficiencies)
        st_entries = [
            {"name": key, "value": value, "proficiency": proficiencies.get(key, 0)}
            for key, value in raw.items()
        ]

    sheets = dict(support_sheets or {})
    sheets.setdefault(sheet_name, matrix)
    if not name and system == "dnd":
        for fallback in ("角色", "起源", "主要", "主要情况"):
            if fallback in sheets:
                name = _extract_name(sheets[fallback], _command_cells(sheets[fallback]))
                if name:
                    break

    if system == "dnd" and not attributes and not skills:
        for fallback in ("主要", "主要情况", "施法"):
            if fallback not in sheets:
                continue
            raw = _extract_dnd_grid(sheets[fallback])
            if raw:
                attributes, skills = _finalize_dnd_values(raw, {})
                st_entries = [{"name": key, "value": value, "proficiency": 0} for key, value in raw.items()]
                break

    if not name:
        raise ValueError(f"已命中工作表“{sheet_name}”，但没有读取到角色姓名")
    if not attributes and not skills:
        raise ValueError(f"已命中工作表“{sheet_name}”，但没有读取到可导入的属性或技能")

    if system == "dnd":
        weapon_matrices = [sheets[name] for name in ("主要", "主要情况") if name in sheets]
        item_matrices = [sheets[name] for name in ("背包",) if name in sheets]
        if not weapon_matrices:
            weapon_matrices = [matrix]
        if not item_matrices:
            item_matrices = [matrix]
        weapons = []
        for source in weapon_matrices:
            weapons = _extract_weapons(source, system, skills)
            if weapons:
                break
        items = []
        for source in item_matrices:
            items = _extract_items(source, system)
            if items:
                break
    else:
        weapons = _extract_weapons(matrix, system, skills)
        items = _extract_items(matrix, system)
    visual_sources = [matrix] + [source for name, source in sheets.items() if source is not matrix]
    appearance = ""
    personality = ""
    for source in visual_sources:
        if not appearance:
            appearance = _extract_labeled_text(source, ("外貌描述", "外貌", "外观描述", "外观", "形象描述", "形象", "外表", "服装", "穿着"))
        if not personality:
            personality = _extract_labeled_text(source, ("性格", "个性", "人格"))
        if appearance and personality:
            break
    card = {
        "name": name,
        "system": system,
        "source_sheet": sheet_name,
        "attributes": attributes,
        "skills": skills,
        "proficiencies": proficiencies,
        "st_entries": st_entries,
        "weapons": weapons,
        "items": items,
        "appearance": appearance,
        "personality": personality,
    }
    if system == "dnd":
        card["spell_slots"] = _extract_spell_slots(sheets)
        card["spell_save_dc"], card["spell_attack_bonus"] = _extract_spellcasting_numbers(sheets, attributes)
        card["spells"] = [_classify_spell(spell) for spell in _extract_selected_spells(sheets)]
    return card


def extract_import_card(file_content: bytes, filename: str, card_system: str = "auto") -> dict[str, Any]:
    requested = card_system if card_system in {"auto", "coc", "dnd"} else "auto"
    sheet_name, detected, matrix, support_sheets = _load_import_matrix(file_content, filename, requested)
    card = parse_import_matrix(matrix, detected, sheet_name, support_sheets)
    card["filename"] = filename
    return card
